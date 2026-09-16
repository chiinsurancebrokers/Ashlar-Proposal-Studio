"""Ashlar-branded client PPTX generator."""
from __future__ import annotations

import io
from datetime import date

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from .client_analysis import plan_display_name, premium_display, client_facing_deductible, find_plan_narrative

NAVY = RGBColor(12, 26, 42)
NAVY2 = RGBColor(18, 36, 58)
INK = RGBColor(28, 37, 55)
MUTED = RGBColor(105, 117, 137)
BORDER = RGBColor(225, 231, 239)
BG = RGBColor(246, 248, 251)
WHITE = RGBColor(255, 255, 255)
PURPLE = RGBColor(108, 76, 245)
GOLD = RGBColor(231, 173, 67)
GREEN = RGBColor(21, 150, 106)
RED = RGBColor(196, 72, 72)
FONT = "Aptos"


def _labels(language: str) -> dict:
    greek = str(language).lower().startswith(("gr", "el")) or "greek" in str(language).lower()
    if greek:
        return {
            "subtitle": "Ανεξάρτητη συγκριτική ανάλυση και αιτιολογημένη συμβουλευτική άποψη",
            "reviewed": "ασφαλιστικές επιλογές εξετάστηκαν",
            "context": "ΠΛΑΙΣΙΟ ΠΕΛΑΤΗ", "what_matters": "Τι έχει σημασία σε αυτή τη σύγκριση",
            "executive": "Συνοπτική εικόνα", "overview": "ΣΥΝΟΨΗ ΕΠΙΛΟΓΩΝ", "glance": "Τα προγράμματα με μια ματιά",
            "premium": "ΑΣΦΑΛΙΣΤΡΟ", "annual": "ΕΤΗΣΙΟ ΟΡΙΟ", "area": "ΠΕΡΙΟΧΗ", "excess": "ΑΠΑΛΛΑΓΗ",
            "plan_review": "ΠΑΡΟΥΣΙΑΣΗ ΠΡΟΓΡΑΜΜΑΤΟΣ", "strengths": "Πλεονεκτήματα", "consider": "Σημεία προσοχής",
            "comparison": "ΣΥΓΚΡΙΤΙΚΗ ΑΠΕΙΚΟΝΙΣΗ", "compare_title": "Πώς συγκρίνονται οι επιλογές", "benefit": "Κάλυψη / όρος",
            "what_matters_sec": "ΟΥΣΙΑΣΤΙΚΕΣ ΔΙΑΦΟΡΕΣ", "diff_title": "Οι διαφορές που έχουν σημασία",
            "assessment": "ΑΞΙΟΛΟΓΗΣΗ ASHLAR", "preferred": "Προτιμώμενη επιλογή", "alternative": "Εναλλακτική",
            "before": "ΠΡΙΝ ΠΡΟΧΩΡΗΣΕΤΕ", "important": "Σημαντικές επισημάνσεις", "next": "Επόμενα βήματα",
            "notice": "Σημαντική σημείωση", "footer": "Ανεξάρτητη συγκριτική ανάλυση",
            "final": "Η τελική επιλογή υπόκειται στο underwriting της ασφαλιστικής και στα ισχύοντα συμβατικά έγγραφα.",
        }
    return {
        "subtitle": "Independent plan comparison and reasoned advisory view", "reviewed": "insurance options reviewed",
        "context": "CLIENT CONTEXT", "what_matters": "What matters in this comparison", "executive": "Executive summary",
        "overview": "OPTIONS OVERVIEW", "glance": "The plans at a glance", "premium": "PREMIUM", "annual": "ANNUAL LIMIT",
        "area": "AREA", "excess": "EXCESS", "plan_review": "PLAN REVIEW", "strengths": "Strengths",
        "consider": "Points to consider", "comparison": "SIDE-BY-SIDE COMPARISON", "compare_title": "How the options compare",
        "benefit": "Benefit / term", "what_matters_sec": "WHAT MATTERS", "diff_title": "The differences that matter",
        "assessment": "ASHLAR ASSESSMENT", "preferred": "Preferred fit", "alternative": "Alternative",
        "before": "BEFORE YOU PROCEED", "important": "Important considerations", "next": "Next steps",
        "notice": "Important notice", "footer": "Independent comparative analysis",
        "final": "The final choice remains subject to insurer underwriting and the governing policy documents.",
    }


def _rect(slide, x, y, w, h, fill, radius=False, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line
    else:
        shp.line.fill.background()
    return shp


def _text(slide, text, x, y, w, h, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True; tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = str(text or "")
    run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    return box


def _bullets(slide, items, x, y, w, h, size=14, color=INK, bullet_color=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    for i, item in enumerate(items or []):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = str(item); p.level = 0; p.space_after = Pt(6)
        p.font.name = FONT; p.font.size = Pt(size); p.font.color.rgb = color
        p.text = "• " + p.text
    return box


def _base(prs, section="ASHLAR ASSURANCE"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = BG
    _rect(slide, 0, 0, 13.333, .09, GOLD)
    _text(slide, section, .55, .26, 4, .3, 9, True, MUTED)
    _text(slide, "Ashlar Assurance", 10.65, .26, 2.1, .3, 9, True, MUTED, PP_ALIGN.RIGHT)
    return slide


def _footer(slide, page_note="Independent comparative analysis"):
    _text(slide, page_note, .55, 7.15, 8.5, .2, 8, False, MUTED)
    _text(slide, date.today().strftime("%d %b %Y"), 11.2, 7.15, 1.55, .2, 8, False, MUTED, PP_ALIGN.RIGHT)


def _compact(value, max_chars=90):
    s = str(value or "Not specified").replace("\n", " ").strip()
    return s if len(s) <= max_chars else s[: max_chars - 1].rstrip() + "…"


def build_pptx_bytes(*, client_analysis: dict, results: list[dict], language: str = "English") -> bytes:
    lab = _labels(language)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
    _rect(s, 0, 0, .12, 7.5, GOLD)
    _text(s, "ASHLAR ASSURANCE", .7, .65, 4, .35, 11, True, GOLD)
    _text(s, client_analysis.get("report_title") or "Insurance Comparative Analysis", .7, 1.45, 11.7, 1.4, 31, True, WHITE)
    _text(s, client_analysis.get("client_name") or "Client", .7, 3.15, 10, .55, 20, True, WHITE)
    _text(s, lab["subtitle"], .7, 3.75, 8.8, .55, 14, False, RGBColor(188, 200, 215))
    _rect(s, .7, 5.45, 4.15, .65, NAVY2, True, RGBColor(52, 74, 98))
    _text(s, f"{len(results)} {lab["reviewed"]}", .95, 5.61, 3.65, .28, 12, True, WHITE)
    _text(s, date.today().strftime("%d %B %Y"), .7, 6.68, 2.6, .25, 9, False, RGBColor(160, 175, 193))

    # Client needs + executive summary
    s = _base(prs, lab["context"])
    _text(s, lab["what_matters"], .55, .8, 12, .55, 25, True)
    _text(s, client_analysis.get("client_needs_summary") or client_analysis.get("client_priorities") or "No specific priorities supplied.", .55, 1.48, 5.9, 2.2, 15, False, INK)
    _rect(s, 6.75, 1.35, 5.95, 4.85, WHITE, True, BORDER)
    _text(s, lab["executive"], 7.05, 1.68, 5.25, .45, 17, True, PURPLE)
    _text(s, client_analysis.get("executive_summary") or "", 7.05, 2.28, 5.25, 3.55, 14, False, INK)
    _footer(s)

    # Overview cards, max 4 per slide
    for start in range(0, len(results), 4):
        subset = results[start:start+4]
        s = _base(prs, lab["overview"])
        _text(s, lab["glance"], .55, .78, 12, .55, 25, True)
        n = len(subset); gap = .22; left = .55; total_w = 12.23; card_w = (total_w - gap*(n-1))/n
        for idx, result in enumerate(subset):
            a = result.get("analysis") or {}; x = left + idx*(card_w+gap)
            _rect(s, x, 1.55, card_w, 4.92, WHITE, True, BORDER)
            _rect(s, x, 1.55, card_w, .08, PURPLE if idx % 2 == 0 else GOLD)
            _text(s, a.get("provider") or result.get("provider"), x+.22, 1.82, card_w-.44, .3, 10, True, MUTED)
            _text(s, a.get("plan_name") or result.get("target_plan") or "Plan", x+.22, 2.17, card_w-.44, .62, 18, True, INK)
            _text(s, lab["premium"], x+.22, 2.95, card_w-.44, .22, 8, True, MUTED)
            _text(s, _compact(premium_display(a), 38), x+.22, 3.20, card_w-.44, .48, 14, True)
            _text(s, lab["annual"], x+.22, 3.82, card_w-.44, .22, 8, True, MUTED)
            _text(s, _compact(a.get("annual_limit"), 38), x+.22, 4.06, card_w-.44, .55, 13, True)
            _text(s, lab["area"], x+.22, 4.73, card_w-.44, .22, 8, True, MUTED)
            _text(s, _compact(a.get("area_of_cover"), 50), x+.22, 4.98, card_w-.44, .65, 12, True)
            _text(s, lab["excess"], x+.22, 5.78, card_w-.44, .22, 8, True, MUTED)
            _text(s, _compact(client_facing_deductible(a.get("deductible_or_excess")), 48), x+.22, 6.02, card_w-.44, .34, 10, False)
        _footer(s)

    # Each plan
    narrative_plans = client_analysis.get("plans", [])
    for result in results:
        a = result.get("analysis") or {}
        n = find_plan_narrative(narrative_plans, result)
        s = _base(prs, lab["plan_review"])
        _text(s, a.get("provider") or result.get("provider"), .55, .75, 4, .35, 11, True, PURPLE)
        _text(s, a.get("plan_name") or result.get("target_plan") or "Plan", .55, 1.12, 8.5, .7, 27, True)
        _text(s, n.get("positioning") or "", .55, 1.82, 11.7, .55, 13, False, MUTED)
        # facts strip
        facts = [
            ("Premium", premium_display(a)),
            ("Annual limit", a.get("annual_limit")),
            ("Excess", client_facing_deductible(a.get("deductible_or_excess"))),
            ("Area", a.get("area_of_cover")),
        ]
        for i, (fact_label, val) in enumerate(facts):
            x = .55 + i*3.05
            _rect(s, x, 2.48, 2.82, 1.1, WHITE, True, BORDER)
            _text(s, fact_label.upper(), x+.18, 2.68, 2.45, .2, 8, True, MUTED)
            _text(s, _compact(val, 44), x+.18, 2.96, 2.45, .45, 11, True)
        _text(s, n.get("summary") or "", .55, 3.95, 5.9, 1.75, 13, False)
        _text(s, lab["strengths"], 6.82, 3.95, 2.3, .35, 15, True, GREEN)
        _bullets(s, (n.get("strengths") or [])[:4], 6.82, 4.38, 5.75, 1.15, 11)
        _text(s, lab["consider"], 6.82, 5.58, 2.7, .35, 15, True, GOLD)
        _bullets(s, (n.get("considerations") or [])[:3], 6.82, 5.98, 5.75, .82, 10)
        _footer(s)

    # Comparison matrix - split rows
    matrix = client_analysis.get("comparison_matrix") or []
    names = [plan_display_name(r) for r in results]
    for start in range(0, len(matrix), 7):
        rows = matrix[start:start+7]
        s = _base(prs, lab["comparison"])
        _text(s, lab["compare_title"], .55, .76, 12, .55, 25, True)
        left = .45; top = 1.55; first_w = 2.15; rem = 12.45-first_w; col_w = rem/max(len(names),1); row_h = .69
        _rect(s, left, top, first_w, row_h, NAVY)
        _text(s, lab["benefit"], left+.12, top+.18, first_w-.24, .25, 9, True, WHITE)
        for j, name in enumerate(names):
            x = left+first_w+j*col_w
            _rect(s, x, top, col_w, row_h, NAVY2)
            _text(s, _compact(name, 32), x+.09, top+.12, col_w-.18, .45, 8, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        for i, row in enumerate(rows):
            y = top+row_h*(i+1); fill = WHITE if i%2==0 else RGBColor(249,250,252)
            _rect(s, left, y, first_w, row_h, fill, False, BORDER)
            _text(s, row.get("topic"), left+.12, y+.12, first_w-.24, .45, 9, True)
            for j, name in enumerate(names):
                x = left+first_w+j*col_w
                _rect(s, x, y, col_w, row_h, fill, False, BORDER)
                _text(s, _compact((row.get("values") or {}).get(name), 54), x+.09, y+.08, col_w-.18, .52, 8, False)
        _footer(s, lab["footer"])

    # Key differences
    diffs = client_analysis.get("key_differences") or []
    if diffs:
        s = _base(prs, lab["what_matters_sec"])
        _text(s, lab["diff_title"], .55, .76, 12, .55, 25, True)
        y = 1.58
        for i, d in enumerate(diffs[:4], 1):
            _rect(s, .55, y, 12.2, 1.14, WHITE, True, BORDER)
            _text(s, f"{i:02d}", .77, y+.25, .45, .3, 11, True, PURPLE)
            _text(s, d.get("title") or "Difference", 1.32, y+.17, 3.15, .34, 14, True)
            _text(s, d.get("analysis") or "", 4.5, y+.13, 5.15, .62, 10, False)
            _text(s, d.get("client_impact") or "", 9.85, y+.13, 2.55, .68, 9, True, MUTED)
            y += 1.28
        _footer(s)

    # Assessment
    ass = client_analysis.get("ashlar_assessment") or {}
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
    _text(s, lab["assessment"], .65, .58, 4, .3, 10, True, GOLD)
    _text(s, ass.get("headline") or "Our view", .65, 1.05, 11.7, .95, 27, True, WHITE)
    if ass.get("recommended_provider") or ass.get("recommended_plan"):
        _rect(s, .65, 2.22, 4.1, .75, PURPLE, True)
        _text(s, f"{lab["preferred"]}: {ass.get('recommended_provider','')} {ass.get('recommended_plan','')}", .9, 2.43, 3.6, .32, 13, True, WHITE)
    _bullets(s, (ass.get("reasoning") or [])[:5], .65, 3.18, 7.45, 2.55, 14, WHITE)
    _rect(s, 8.45, 2.22, 4.2, 3.7, NAVY2, True, RGBColor(53,75,98))
    _text(s, lab["alternative"], 8.75, 2.55, 3.5, .35, 13, True, GOLD)
    _text(s, f"{ass.get('alternative_provider','')} {ass.get('alternative_plan','')}", 8.75, 3.0, 3.5, .52, 18, True, WHITE)
    _text(s, ass.get("alternative_reason") or "", 8.75, 3.7, 3.45, 1.1, 11, False, RGBColor(204,214,225))
    _text(s, ass.get("when_the_alternative_may_be_better") or "", 8.75, 4.92, 3.45, .7, 10, False, RGBColor(173,188,204))
    _text(s, lab["final"], .65, 6.72, 11.8, .3, 9, False, RGBColor(157,174,193))

    # Considerations / next steps
    s = _base(prs, lab["before"])
    _text(s, lab["important"], .55, .78, 5.9, .5, 23, True)
    _bullets(s, (client_analysis.get("important_considerations") or [])[:7], .55, 1.55, 5.8, 4.55, 12)
    _rect(s, 6.72, 1.42, 5.95, 4.9, WHITE, True, BORDER)
    _text(s, lab["next"], 7.02, 1.78, 5.2, .45, 18, True, PURPLE)
    _bullets(s, (client_analysis.get("next_steps") or [])[:6], 7.02, 2.42, 5.1, 2.55, 12)
    _text(s, lab["notice"], 7.02, 5.18, 2.3, .28, 10, True, GOLD)
    _text(s, client_analysis.get("disclaimer") or "", 7.02, 5.55, 5.1, .62, 9, False, MUTED)
    _footer(s)

    out = io.BytesIO(); prs.save(out); return out.getvalue()
