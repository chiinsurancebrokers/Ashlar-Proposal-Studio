"""Excel and Markdown exports.

Formatting only, so the tests are about what must not go wrong: a verdict that
no longer belongs to its row, a document sentence turning into a spreadsheet
formula, a clause cut from the client report disappearing from the export.
"""
from __future__ import annotations

import io

from openpyxl import load_workbook

from core.export_tables import (
    STATUS_CONFIRMED, STATUS_DISPUTED, STATUS_MISSED, STATUS_STALE, STATUS_UNCHECKED,
    comparison_table, limitation_table, to_excel_bytes, to_markdown,
)
from core.limitation_audit import (
    AGREES, DISAGREES, LimitationAuditResult, LimitationFinding, MissedLimitation,
)

PLAN = "APRIL - Essential"


def _result(limitations, provider="APRIL", plan="Essential"):
    return {
        "provider": provider, "target_plan": plan,
        "analysis": {
            "provider": provider, "plan_name": plan,
            "premium": {"amount": "1200", "currency": "EUR", "frequency": "per year"},
            "deductible_or_excess": "EUR 500",
            "annual_limit": "EUR 1,000,000",
            "area_of_cover": "Worldwide excluding USA",
            "underwriting": {"basis": "Moratorium", "pre_existing_conditions": "Excluded for 24 months"},
            "waiting_periods": [{"topic": "Maternity", "detail": "12 months"}, "Dental 3 months"],
            "benefits": {"inpatient": "Paid in full", "dental": "EUR 1,000 | 80%"},
            "critical_limitations": limitations,
        },
    }


LIMITS = [
    {"topic": "Ventilation", "detail": "90 days maximum for life support.", "source_page": 41},
    {"topic": "Physio", "detail": "Maximum 10 sessions per year."},
]


def _finding(key, detail, verdict, note=""):
    return LimitationFinding(plan=PLAN, key=key, topic="t", detail=detail, verdict=verdict, note=note)


def _audit(findings, missed=()):
    return LimitationAuditResult(findings=list(findings), missed=list(missed), ran=True, model="gpt-4o")


# ---------------------------------------------------------------------------
# Limitations table
# ---------------------------------------------------------------------------


def test_every_limitation_is_exported_even_without_a_second_opinion():
    header, rows = limitation_table([_result(LIMITS)], None)
    assert len(rows) == 2
    assert {r[-1] for r in rows} == {STATUS_UNCHECKED}
    assert "Source page" in header


def test_a_verdict_is_shown_beside_its_own_limitation():
    audit = _audit([
        _finding("L1", LIMITS[0]["detail"], AGREES),
        _finding("L2", LIMITS[1]["detail"], DISAGREES, "10 sessions per policy year, referral needed"),
    ])
    _, rows = limitation_table([_result(LIMITS)], audit)
    assert rows[0][-1] == STATUS_CONFIRMED
    assert rows[1][-1] == STATUS_DISPUTED
    assert "referral needed" in rows[1][5]


def test_a_verdict_is_not_carried_over_when_the_analysis_has_changed():
    audit = _audit([_finding("L1", "30 days maximum for life support.", AGREES)])
    _, rows = limitation_table([_result(LIMITS)], audit)
    assert rows[0][-1] == STATUS_STALE
    assert rows[0][4] == ""


def test_a_check_that_did_not_run_gives_no_verdicts():
    failed = LimitationAuditResult(ran=False, error="no key")
    _, rows = limitation_table([_result(LIMITS)], failed)
    assert {r[-1] for r in rows} == {STATUS_UNCHECKED}


def test_clauses_the_second_reader_says_were_missed_are_listed_for_the_broker():
    audit = _audit([_finding("L1", LIMITS[0]["detail"], AGREES)],
                   [MissedLimitation(plan=PLAN, topic="Sport", quote="any professional sport")])
    _, rows = limitation_table([_result(LIMITS)], audit)
    missed = [r for r in rows if r[-1] == STATUS_MISSED]
    assert len(missed) == 1 and missed[0][5] == "any professional sport"


# ---------------------------------------------------------------------------
# Comparison table
# ---------------------------------------------------------------------------


def test_comparison_has_one_column_per_plan():
    header, rows = comparison_table([_result(LIMITS), _result([], "Bupa", "Select")])
    assert header == ["Item", "APRIL - Essential", "Bupa - Select"]
    assert all(len(r) == 3 for r in rows)
    by_item = {r[0]: r for r in rows}
    assert by_item["Annual limit"][1] == "EUR 1,000,000"
    assert by_item["Waiting periods"][1] == "12 months; Dental 3 months"


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def _workbook(results, audit=None):
    return load_workbook(io.BytesIO(to_excel_bytes(results, audit)))


def test_the_workbook_has_the_three_sheets():
    wb = _workbook([_result(LIMITS)], _audit([_finding("L1", LIMITS[0]["detail"], AGREES)]))
    assert wb.sheetnames == ["Comparison", "Limitations", "About"]
    assert wb["Limitations"].max_row == 3  # header + both limitations
    assert wb["Limitations"].cell(row=2, column=7).value == STATUS_CONFIRMED


def test_document_text_that_looks_like_a_formula_stays_text():
    hostile = [{"topic": "=HYPERLINK(\"http://example.com\",\"x\")", "detail": "=1+1 and 90 days"}]
    wb = _workbook([_result(hostile)])
    cell = wb["Limitations"].cell(row=2, column=3)
    assert cell.data_type != "f"
    assert cell.value == "=1+1 and 90 days"
    assert wb["Limitations"].cell(row=2, column=2).data_type != "f"


def test_the_about_sheet_says_whether_the_second_reader_ran():
    assert "not run" in _workbook([_result(LIMITS)])["About"]["A3"].value
    ran = _workbook([_result(LIMITS)], _audit([_finding("L1", LIMITS[0]["detail"], AGREES)]))["About"]["A3"].value
    assert "1 confirmed" in ran and "gpt-4o" in ran


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------


def test_markdown_escapes_pipes_so_a_table_row_cannot_break():
    md = to_markdown([_result(LIMITS)])
    assert "EUR 1,000 \\| 80%" in md
    for line in md.splitlines():
        if line.startswith("| ") and "Dental" in line:
            assert line.count("|") - line.count("\\|") == 3  # item | one plan | closing


def test_markdown_has_the_limitations_section_with_statuses():
    md = to_markdown([_result(LIMITS)], _audit([_finding("L1", LIMITS[0]["detail"], DISAGREES, "60 days")]))
    assert "## Limitations" in md
    assert STATUS_DISPUTED in md


def test_markdown_says_so_when_there_are_no_limitations():
    assert "No limitations were extracted." in to_markdown([_result([])])
