"""Tabular exports of a case: Excel and Markdown.

Two views of the same facts:

* Comparison: one column per plan, one row per item (premium, limits, benefit
  lines). Built for reading across plans.
* Limitations: one row per coverage-limiting clause, with the second reader's
  verdict beside it. The client report keeps only the few most material
  limitations; this is the full list, so a clause cut from the report is still
  somewhere the broker can see it.

These are formatting only. They make no model call.

A verdict is shown only when it belongs to the text beside it. If the analysis
changed after the second-opinion check ran, the old verdict is not carried
across: the row says it was not checked.
"""
from __future__ import annotations

import io
import re
from typing import Any

from .limitation_audit import AGREES, DISAGREES, LimitationAuditResult, limitation_lines

COMPARISON_SHEET = "Comparison"
LIMITATIONS_SHEET = "Limitations"
ABOUT_SHEET = "About"

LIMITATION_HEADERS = [
    "Provider / plan", "Topic", "Limitation", "Source page",
    "Second-opinion check", "What the document says", "Status",
]

STATUS_CONFIRMED = "Confirmed by second reader"
STATUS_DISPUTED = "DISPUTED - second reader disagrees"
STATUS_UNCONFIRMED = "Not confirmed - document silent or figure not found"
STATUS_UNCHECKED = "Not checked"
STATUS_STALE = "Not checked - analysis changed after the check"
STATUS_MISSED = "Possibly missed - broker to decide"


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dict):
        return str(value.get("detail") or value.get("topic") or "").strip()
    if isinstance(value, (list, tuple)):
        return "; ".join(t for t in (_text(v) for v in value) if t)
    return str(value).strip()


# ---------------------------------------------------------------------------
# Rows
# ---------------------------------------------------------------------------

def comparison_table(results: list[dict]) -> tuple[list[str], list[list[str]]]:
    """Header plus rows, one column per plan."""
    from .client_analysis import BENEFIT_ORDER, plan_display_name, premium_display

    names = [plan_display_name(r) for r in results or []]
    analyses = [(r.get("analysis") or {}) for r in results or []]

    def underwriting(a: dict, key: str) -> str:
        return _text((a.get("underwriting") or {}).get(key))

    items: list[tuple[str, list[str]]] = [
        ("Premium", [premium_display(a) for a in analyses]),
        ("Deductible / excess", [_text(a.get("deductible_or_excess")) for a in analyses]),
        ("Annual limit", [_text(a.get("annual_limit")) for a in analyses]),
        ("Area of cover", [_text(a.get("area_of_cover")) for a in analyses]),
        ("Underwriting basis", [underwriting(a, "basis") for a in analyses]),
        ("Pre-existing conditions", [underwriting(a, "pre_existing_conditions") for a in analyses]),
        ("Waiting periods", [_text(a.get("waiting_periods")) for a in analyses]),
    ]
    for key, label in BENEFIT_ORDER:
        items.append((label, [_text((a.get("benefits") or {}).get(key)) for a in analyses]))

    header = ["Item", *names]
    rows = [[label, *values] for label, values in items]
    return header, rows


def _status(finding, detail: str) -> tuple[str, str, str]:
    """(check label, what the document says, status) for one limitation row."""
    if finding is None:
        return "", "", STATUS_UNCHECKED
    if finding.detail.strip() != detail.strip():
        return "", "", STATUS_STALE
    if finding.verdict == AGREES:
        return "agrees", finding.note, STATUS_CONFIRMED
    if finding.verdict == DISAGREES:
        return "disagrees", finding.note, STATUS_DISPUTED
    return "not found", finding.note, STATUS_UNCONFIRMED


def limitation_table(results: list[dict], audit: LimitationAuditResult | None = None
                     ) -> tuple[list[str], list[list[str]]]:
    """Every limitation of every plan, with the second reader's verdict if it ran."""
    from .client_analysis import plan_display_name

    by_key = {}
    if audit is not None and audit.ran:
        by_key = {(f.plan, f.key): f for f in audit.findings}

    rows: list[list[str]] = []
    for result in results or []:
        plan = plan_display_name(result)
        for key, topic, detail, page in limitation_lines(result.get("analysis") or {}):
            check, says, status = _status(by_key.get((plan, key)), detail)
            rows.append([plan, topic, detail, page, check, says, status])
        if audit is not None and audit.ran:
            for missed in audit.missed:
                if missed.plan == plan:
                    rows.append([plan, missed.topic, "", "", "possibly missed",
                                 missed.quote, STATUS_MISSED])
    return LIMITATION_HEADERS, rows


def about_lines(audit: LimitationAuditResult | None) -> list[str]:
    lines = [
        "Facts are extracted by AI from the provider documents and the client's quotation.",
        "Check them against the source documents before presenting them to a client.",
    ]
    if audit is None:
        lines.append("Second-opinion check on the limitations: not run for this export.")
    elif not audit.ran:
        lines.append("Second-opinion check on the limitations: did not complete. " + (audit.error or ""))
    else:
        counts = {
            "confirmed": sum(1 for f in audit.findings if f.verdict == AGREES),
            "disputed": len(audit.disagreements),
            "not confirmed": len(audit.unverified),
            "possibly missed": len(audit.missed),
        }
        lines.append(
            f"Second-opinion check on the limitations ({audit.model or 'second reader'}): "
            + ", ".join(f"{n} {label}" for label, n in counts.items()) + "."
        )
    return lines


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def _put(ws, row: int, col: int, value: str):
    """Write text, never a formula: document text that starts with '=' stays text."""
    cell = ws.cell(row=row, column=col)
    cell.value = value
    if cell.data_type == "f":
        cell.data_type = "s"
    return cell


def to_excel_bytes(results: list[dict], audit: LimitationAuditResult | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill("solid", fgColor="1F2A44")
    header_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")
    status_fill = {
        STATUS_CONFIRMED: PatternFill("solid", fgColor="DCEFE0"),
        STATUS_DISPUTED: PatternFill("solid", fgColor="F6CFCF"),
        STATUS_MISSED: PatternFill("solid", fgColor="FBE7B8"),
        STATUS_UNCONFIRMED: PatternFill("solid", fgColor="FFF4D1"),
    }

    wb = Workbook()
    ws = wb.active
    ws.title = COMPARISON_SHEET
    header, rows = comparison_table(results)
    for c, value in enumerate(header, start=1):
        cell = _put(ws, 1, c, value)
        cell.fill, cell.font, cell.alignment = header_fill, header_font, wrap
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = _put(ws, r, c, value)
            cell.alignment = wrap
            if c == 1:
                cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 28
    for c in range(2, len(header) + 1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 46
    ws.freeze_panes = "B2"

    lim = wb.create_sheet(LIMITATIONS_SHEET)
    header, rows = limitation_table(results, audit)
    for c, value in enumerate(header, start=1):
        cell = _put(lim, 1, c, value)
        cell.fill, cell.font, cell.alignment = header_fill, header_font, wrap
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = _put(lim, r, c, str(value))
            cell.alignment = wrap
        fill = status_fill.get(row[-1])
        if fill is not None:
            lim.cell(row=r, column=len(row)).fill = fill
    for letter, width in zip("ABCDEFG", (30, 26, 56, 12, 20, 56, 34)):
        lim.column_dimensions[letter].width = width
    lim.freeze_panes = "A2"
    if rows:
        lim.auto_filter.ref = f"A1:G{len(rows) + 1}"

    about = wb.create_sheet(ABOUT_SHEET)
    for r, line in enumerate(about_lines(audit), start=1):
        _put(about, r, 1, line).alignment = wrap
    about.column_dimensions["A"].width = 110

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------

def _cell(value: Any) -> str:
    text = re.sub(r"\s*\n\s*", " ", str(value or "")).strip()
    return text.replace("|", "\\|") or "—"


def _md_table(header: list[str], rows: list[list[str]]) -> str:
    out = ["| " + " | ".join(_cell(h) for h in header) + " |",
           "| " + " | ".join("---" for _ in header) + " |"]
    out += ["| " + " | ".join(_cell(v) for v in row) + " |" for row in rows]
    return "\n".join(out)


def to_markdown(results: list[dict], audit: LimitationAuditResult | None = None,
                *, title: str = "Insurance case comparison") -> str:
    comparison_header, comparison_rows = comparison_table(results)
    limitation_header, limitation_rows = limitation_table(results, audit)
    parts = [
        f"# {title}",
        "## Comparison",
        _md_table(comparison_header, comparison_rows),
        "## Limitations",
        _md_table(limitation_header, limitation_rows) if limitation_rows else "No limitations were extracted.",
        "## About this export",
        "\n".join(f"- {line}" for line in about_lines(audit)),
    ]
    return "\n\n".join(parts) + "\n"
